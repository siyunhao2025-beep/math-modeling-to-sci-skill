#!/usr/bin/env python3
"""Executable quality-gate engine for S1-S6.

The YAML file remains the source of truth for thresholds and severity.
This module evaluates every configured condition by check-name and writes a
machine-readable gate result. Semantic checks use conservative heuristics;
if a condition cannot be evaluated from the available artifacts, it is
reported as ``not_evaluable`` instead of being silently treated as PASS.
"""

from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
import glob
import json
import os
import re
import shutil
import subprocess
from typing import Any, Callable

import common

GATES_YAML = common.repo_path("config", "quality-gates.yaml")

_NUMBER_RE = re.compile(
    r"(?<![A-Za-z_])[-+]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][-+]?\d+)?"
)
_RESEARCH_GAP_RE = re.compile(
    r"\b(gap|remain(?:s|ed)?|however|limited|lack(?:s|ed)?|unresolved|"
    r"not yet|few studies|little is known)\b|研究空白|尚未|缺乏|不足",
    re.I,
)
_CONTRIB_RE = re.compile(
    r"\b(contribution|we propose|we develop|we introduce|this (?:study|work) "
    r"(?:proposes|develops|introduces|contributes))\b|本文(?:提出|构建|发展|贡献)",
    re.I,
)
_EXPERIMENT_CLAIM_RE = re.compile(
    r"\b(?:experiments?\s+(?:show|demonstrate|confirm)|results?\s+indicate|"
    r"we\s+achieved\b.*\b(?:accuracy|improvement|speedup))",
    re.I,
)
_MATH_MARKER_RE = re.compile(
    r"\$\$|(?<!\\)\$[^$\n]+\$|\\\(|\\\[|\\begin\{(?:equation|align|gather|multline)"
)


@dataclass
class ConditionResult:
    id: str
    check: str
    severity: str
    status: str
    detail: str

    def to_dict(self) -> dict[str, Any]:
        return {
            "id": self.id,
            "check": self.check,
            "severity": self.severity,
            "status": self.status,
            "detail": self.detail,
        }


def _load_json(path: str) -> dict:
    if not os.path.isfile(path):
        return {}
    return common.load_json(path)


def _all_text(ms: dict) -> str:
    parts: list[str] = []
    meta = ms.get("meta", {})
    if meta.get("title"):
        parts.append(str(meta["title"]))
    if meta.get("abstract"):
        parts.append(str(meta["abstract"]))
    for sec in ms.get("sections", []):
        if sec.get("heading"):
            parts.append(str(sec["heading"]))
        for block in sec.get("blocks", []):
            if block.get("text"):
                parts.append(str(block["text"]))
            if block.get("latex"):
                parts.append(str(block["latex"]))
    return "\n".join(parts)


def _normalized_source_text(path: str | None) -> str:
    if not path or not os.path.isfile(path):
        return ""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        try:
            from docx import Document
        except Exception:
            return ""
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells if cell.text.strip())
        return "\n".join(parts)
    try:
        text = open(path, encoding="utf-8").read()
    except Exception:
        return ""
    if ext == ".tex":
        text = re.sub(r"(?m)^\s*%.*$", "", text)
        text = re.sub(r"\\(?:documentclass|usepackage)(?:\[[^\]]*\])?\{[^}]*\}", " ", text)
        text = re.sub(r"\\(?:begin|end)\{[^}]*\}", " ", text)
        text = re.sub(r"\\(?:section|subsection|subsubsection|title)\*?\{([^}]*)\}", r" \1 ", text)
        text = re.sub(r"\\[A-Za-z@]+(?:\[[^\]]*\])?", " ", text)
        text = text.replace("{", " ").replace("}", " ")
    return re.sub(r"\s+", " ", text).strip()


def _number_counter(ms: dict) -> Counter[str]:
    return Counter(m.group(0) for m in _NUMBER_RE.finditer(_all_text(ms)))


def _dimension_score(value: Any) -> float | None:
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, dict) and isinstance(value.get("score"), (int, float)):
        return float(value["score"])
    return None


def _schema_valid(workdir: str, condition: dict) -> tuple[bool | None, str]:
    target = condition.get("target")
    against = condition.get("against")
    if not target or not against:
        return None, "schema target/against missing in gate config"
    data_path = os.path.join(workdir, target)
    schema_path = common.repo_path(*against.split("/"))
    if not os.path.isfile(data_path):
        return False, f"artifact missing: {target}"
    data = common.load_json(data_path)
    errors = common.validate_against_schema(data, schema_path)
    return (not errors), ("schema valid" if not errors else "; ".join(errors[:5]))


class GateEvaluator:
    def __init__(self, workdir: str, source_path: str | None = None):
        self.workdir = workdir
        self.source_path = source_path
        self.cfg = common.load_yaml(GATES_YAML)
        self.handlers: dict[str, Callable[[dict], tuple[bool | None, str]]] = {
            "schema_valid": self.schema_valid,
            "min_sections": self.min_sections,
            "required_semantic_roles_present": self.required_semantic_roles_present,
            "no_content_loss": self.no_content_loss,
            "math_objects_extracted": self.math_objects_extracted,
            "role_confidence": self.role_confidence,
            "figures_have_source": self.figures_have_source,
            "no_unverified_added_references": self.no_unverified_added_references,
            "numeric_consistency": self.numeric_consistency,
            "math_semantics_preserved": self.math_semantics_preserved,
            "no_fabricated_experiment_claims": self.no_fabricated_experiment_claims,
            "terminology_consistency": self.terminology_consistency,
            "gaps_registered": self.gaps_registered,
            "language_unified": self.language_unified,
            "has_research_gap_statement": self.has_research_gap_statement,
            "has_explicit_contributions": self.has_explicit_contributions,
            "min_total_score": self.min_total_score,
            "min_dimension_score": self.min_dimension_score,
            "no_blocker_weakness": self.no_blocker_weakness,
            "tier_not_insufficient": self.tier_not_insufficient,
            "self_check_all_true": self.self_check_all_true,
            "min_confidence": self.min_confidence,
            "min_recommendations": self.min_recommendations,
            "all_have_fit_evidence": self.all_have_fit_evidence,
            "all_have_rejection_risks": self.all_have_rejection_risks,
            "impact_factor_provenance": self.impact_factor_provenance,
            "no_predatory_journals": self.no_predatory_journals,
            "tier_alignment": self.tier_alignment,
            "has_gradient_strategy": self.has_gradient_strategy,
            "template_files_present": self.template_files_present,
            "build_main_exists": self.build_main_exists,
            "all_ir_sections_mapped": self.all_ir_sections_mapped,
            "all_math_objects_rendered": self.all_math_objects_rendered,
            "bib_present_and_nonempty": self.bib_present_and_nonempty,
            "template_source_official": self.template_source_official,
            "zero_errors": self.zero_errors,
            "citation_bidirectional_clean": self.citation_bidirectional_clean,
            "numeric_consistency_clean": self.numeric_consistency_clean,
            "no_cjk_residue": self.no_cjk_residue,
            "no_placeholder_residue": self.no_placeholder_residue,
            "latex_compiles": self.latex_compiles,
            "journal_constraints_satisfied": self.journal_constraints_satisfied,
            "no_orphan_numbered_objects": self.no_orphan_numbered_objects,
            "language_quality": self.language_quality,
        }

    @property
    def ir(self) -> dict:
        return _load_json(os.path.join(self.workdir, "01-parse", "manuscript.ir.json"))

    @property
    def rewritten(self) -> dict:
        return _load_json(os.path.join(self.workdir, "02-rewrite", "manuscript.rewritten.json"))

    @property
    def assessment(self) -> dict:
        return _load_json(os.path.join(self.workdir, "03-assess", "assessment.json"))

    @property
    def journal_match(self) -> dict:
        return _load_json(os.path.join(self.workdir, "04-journals", "journal-match.json"))

    @property
    def manifest(self) -> dict:
        return _load_json(os.path.join(self.workdir, "05-template", "MANIFEST.json"))

    @property
    def validation(self) -> dict:
        return _load_json(os.path.join(self.workdir, "06-validate", "validation-final.json"))

    def evaluate(self, gate_id: str) -> dict[str, Any]:
        gate = (self.cfg.get("gates") or {}).get(gate_id)
        if not gate:
            raise KeyError(f"gate not found in config: {gate_id}")
        results: list[ConditionResult] = []
        for cond in gate.get("conditions", []):
            check = cond.get("check", "")
            handler = self.handlers.get(check)
            if handler is None:
                ok, detail = None, f"unsupported check: {check}"
            else:
                try:
                    ok, detail = handler(cond)
                except Exception as exc:
                    ok, detail = False, f"check raised {type(exc).__name__}: {exc}"
            if ok is None and cond.get("skippable_when"):
                status = "skipped"
            else:
                status = "not_evaluable" if ok is None else ("pass" if ok else "fail")
            results.append(
                ConditionResult(
                    id=cond.get("id", check),
                    check=check,
                    severity=cond.get("severity", "error"),
                    status=status,
                    detail=detail,
                )
            )
        error_failures = [
            r for r in results
            if r.severity == "error" and r.status in ("fail", "not_evaluable")
        ]
        warning_failures = [
            r for r in results
            if r.severity == "warning" and r.status != "pass"
        ]
        on_fail = (gate.get("on_fail") or {}).get("error") or {}
        result = {
            "gate": gate_id,
            "name": gate.get("name"),
            "stage": gate.get("stage"),
            "hard": bool(gate.get("hard", False)),
            "passed": not error_failures,
            "error_failures": [r.to_dict() for r in error_failures],
            "warnings": [r.to_dict() for r in warning_failures],
            "conditions": [r.to_dict() for r in results],
            "on_fail": on_fail,
        }
        out_dir = os.path.join(self.workdir, "gate-results")
        os.makedirs(out_dir, exist_ok=True)
        common.save_json_atomic(os.path.join(out_dir, f"{gate_id}.json"), result)
        return result

    def schema_valid(self, c): return _schema_valid(self.workdir, c)

    def min_sections(self, c):
        n = len(self.ir.get("sections", []))
        threshold = int(c.get("threshold", 0))
        return n >= threshold, f"sections={n}, threshold={threshold}"

    def required_semantic_roles_present(self, c):
        roles = {s.get("semantic_role") for s in self.ir.get("sections", [])}
        missing = []
        for group in c.get("required_any_of", []):
            if not any(x in roles for x in group):
                missing.append(group)
        return not missing, f"missing role groups={missing}" if missing else "all role groups present"

    def no_content_loss(self, c):
        source = _normalized_source_text(self.source_path)
        if not source:
            return None, "source text unavailable"
        target = _all_text(self.ir)
        source_len = max(1, len(re.sub(r"\s+", "", source)))
        target_len = len(re.sub(r"\s+", "", target))
        ratio = target_len / source_len
        minimum = float(c.get("min_ratio", 0.85))
        return ratio >= minimum, f"normalized text ratio={ratio:.3f}, min={minimum:.3f}"

    def math_objects_extracted(self, c):
        source = _normalized_source_text(self.source_path)
        if not source:
            return None, "source text unavailable"
        has_markers = bool(_MATH_MARKER_RE.search(source))
        count = len(self.ir.get("equations", []))
        return (not has_markers or count > 0), f"source_has_math={has_markers}, equations={count}"

    def role_confidence(self, c):
        values = [
            s.get("role_confidence")
            for s in self.ir.get("sections", [])
            if isinstance(s.get("role_confidence"), (int, float))
        ]
        if not values:
            return None, "no role_confidence values"
        avg = sum(values) / len(values)
        minimum = float(c.get("min_avg", 0.0))
        return avg >= minimum, f"avg role confidence={avg:.3f}, min={minimum:.3f}"

    def figures_have_source(self, c):
        missing = [f.get("id") for f in self.ir.get("figures", []) if not f.get("path")]
        return not missing, f"missing figure source={missing}" if missing else "all figures have source paths"

    def no_unverified_added_references(self, c):
        bad = []
        for ref in self.rewritten.get("references", []):
            if ref.get("origin") == "added_by_s2":
                status = str((ref.get("verification") or {}).get("status", ""))
                if not status.startswith("verified_"):
                    bad.append(ref.get("key") or ref.get("title"))
        return not bad, f"unverified added refs={bad}" if bad else "all added references verified"

    def numeric_consistency(self, c):
        if not self.ir or not self.rewritten:
            return None, "original or rewritten IR missing"
        a, b = _number_counter(self.ir), _number_counter(self.rewritten)
        if a == b:
            return True, f"numeric token multiset preserved ({sum(a.values())} tokens)"
        removed = list((a - b).elements())[:10]
        added = list((b - a).elements())[:10]
        return False, f"numbers changed; removed={removed}, added={added}"

    def math_semantics_preserved(self, c):
        original = {e.get("id"): re.sub(r"\s+", "", e.get("latex", "")) for e in self.ir.get("equations", [])}
        rewritten = {e.get("id"): re.sub(r"\s+", "", e.get("latex", "")) for e in self.rewritten.get("equations", [])}
        missing = [eid for eid in original if eid not in rewritten]
        changed = [eid for eid, value in original.items() if eid in rewritten and rewritten[eid] != value]
        ok = not missing and not changed and len(rewritten) >= len(original)
        return ok, f"missing={missing}, changed={changed}, counts={len(original)}->{len(rewritten)}"

    def no_fabricated_experiment_claims(self, c):
        text = _all_text(self.rewritten)
        if not _EXPERIMENT_CLAIM_RE.search(text):
            return True, "no high-risk experiment-claim pattern found"
        roles = {s.get("semantic_role") for s in self.rewritten.get("sections", []) if s.get("blocks")}
        backed = bool({"experiments", "results"} & roles)
        return backed, f"experiment/result claim found; backing roles present={backed}"

    def terminology_consistency(self, c):
        sm = self.rewritten.get("symbol_map", {}) or {}
        seen: dict[str, str] = {}
        conflicts = []
        for new, info in sm.items():
            old = (info or {}).get("renamed_from")
            if not old:
                continue
            if old in seen and seen[old] != new:
                conflicts.append((old, seen[old], new))
            seen[old] = new
        return not conflicts, f"rename conflicts={conflicts}" if conflicts else "symbol rename mapping consistent"

    def gaps_registered(self, c):
        missing_flags = 0
        for sec in self.rewritten.get("sections", []):
            for block in sec.get("blocks", []):
                if "MISSING" in (block.get("flags") or []):
                    missing_flags += 1
                missing_flags += str(block.get("text") or "").count("[[MISSING")
        gaps = self.rewritten.get("gaps", []) or []
        return len(gaps) >= missing_flags, f"missing markers={missing_flags}, gaps={len(gaps)}"

    def language_unified(self, c):
        if (self.rewritten.get("meta") or {}).get("target_language", "en") != "en":
            return True, "target language is not en"
        hits = common.find_cjk(_all_text(self.rewritten))
        return not hits, f"CJK chars={len(hits)}"

    def has_research_gap_statement(self, c):
        found = bool(_RESEARCH_GAP_RE.search(_all_text(self.rewritten)))
        return found, f"research-gap cue found={found}"

    def has_explicit_contributions(self, c):
        found = bool(_CONTRIB_RE.search(_all_text(self.rewritten)))
        return found, f"contribution cue found={found}"

    def min_total_score(self, c):
        score = self.assessment.get("total_score")
        if not isinstance(score, (int, float)):
            return False, "total_score missing"
        threshold = float(c.get("threshold", 0))
        return score >= threshold, f"total_score={score}, threshold={threshold}"

    def min_dimension_score(self, c):
        dims = self.assessment.get("dimensions") or {}
        failed = {}
        for name, threshold in (c.get("thresholds") or {}).items():
            score = _dimension_score(dims.get(name))
            if score is None or score < float(threshold):
                failed[name] = {"score": score, "threshold": threshold}
        return not failed, f"dimension failures={failed}" if failed else "all dimensions above threshold"

    def no_blocker_weakness(self, c):
        blockers = [w for w in self.assessment.get("weaknesses", []) if w.get("severity") == "blocker"]
        return not blockers, f"blocker weaknesses={len(blockers)}"

    def tier_not_insufficient(self, c):
        tier = self.assessment.get("tier")
        return tier != "insufficient-content", f"tier={tier}"

    def self_check_all_true(self, c):
        sc = self.assessment.get("self_check")
        if not isinstance(sc, dict) or not sc:
            return False, "self_check missing"
        bool_values = [v for v in sc.values() if isinstance(v, bool)]
        return bool_values and all(bool_values), f"boolean self-checks={bool_values}"

    def min_confidence(self, c):
        value = self.assessment.get("confidence")
        if not isinstance(value, (int, float)):
            return None, "confidence not provided"
        threshold = float(c.get("threshold", 0))
        return value >= threshold, f"confidence={value}, threshold={threshold}"

    def min_recommendations(self, c):
        n = len(self.journal_match.get("recommendations", []))
        threshold = int(c.get("threshold", 0))
        return n >= threshold, f"recommendations={n}, threshold={threshold}"

    def all_have_fit_evidence(self, c):
        recs = self.journal_match.get("recommendations", [])
        bad = [r.get("name") for r in recs if not r.get("fit_evidence")]
        return bool(recs) and not bad, f"missing fit evidence={bad}"

    def all_have_rejection_risks(self, c):
        recs = self.journal_match.get("recommendations", [])
        bad = [r.get("name") for r in recs if not r.get("rejection_risks")]
        return bool(recs) and not bad, f"missing rejection risks={bad}"

    def impact_factor_provenance(self, c):
        bad = []
        for r in self.journal_match.get("recommendations", []):
            info = r.get("impact_factor") or {}
            if info.get("value") is not None and not (info.get("year") and info.get("source")):
                bad.append(r.get("name"))
        return not bad, f"IF provenance failures={bad}"

    def no_predatory_journals(self, c):
        value = (self.journal_match.get("self_check") or {}).get("no_predatory_journals")
        return (bool(value) if value is not None else None), f"self_check.no_predatory_journals={value}"

    def tier_alignment(self, c):
        recs = self.journal_match.get("recommendations", [])
        score = (self.journal_match.get("input_summary") or {}).get("total_score")
        if not recs or not isinstance(score, (int, float)):
            return None, "recommendations or input score unavailable"
        if score < 6.0 and recs[0].get("quartile") == "Q1":
            return False, "low-scoring manuscript has Q1 as top recommendation"
        if score >= 8.0 and not any(r.get("quartile") in ("Q1", "Q2") for r in recs):
            return False, "high-scoring manuscript lacks Q1/Q2 option"
        return True, "tier alignment heuristic passed"

    def has_gradient_strategy(self, c):
        note = str(self.journal_match.get("strategy_note") or "")
        found = all(x in note for x in ("冲刺", "稳妥", "保底"))
        return found, f"gradient labels present={found}"

    def template_files_present(self, c):
        patterns = c.get("require_any_of") or []
        found = []
        for pattern in patterns:
            found.extend(glob.glob(os.path.join(self.workdir, pattern)))
        mapping_manifest = os.path.join(self.workdir, "05-template", "MANIFEST.json")
        if os.path.isfile(mapping_manifest):
            found.append(mapping_manifest)
        return bool(found), f"template evidence={found[:5]}"

    def build_main_exists(self, c):
        found = []
        for pattern in c.get("require_any_of", []):
            found.extend(glob.glob(os.path.join(self.workdir, pattern)))
        return bool(found), f"build main={found}"

    def all_ir_sections_mapped(self, c):
        section_ids = {s.get("id") for s in self.rewritten.get("sections", [])}
        mapped = {x.get("ir_id") for x in self.manifest.get("section_map", []) if x.get("placed")}
        missing = sorted(x for x in section_ids - mapped if x)
        return not missing, f"unmapped sections={missing}"

    def all_math_objects_rendered(self, c):
        main_path = os.path.join(self.workdir, "05-template", "build", "main.tex")
        if not os.path.isfile(main_path):
            return False, "main.tex missing"
        text = open(main_path, encoding="utf-8").read()
        eq_expected = len(self.rewritten.get("equations", []))
        fig_expected = len(self.rewritten.get("figures", []))
        tab_expected = len(self.rewritten.get("tables", []))
        eq_actual = text.count(r"\begin{equation}")
        fig_actual = text.count(r"\begin{figure}")
        tab_actual = text.count(r"\begin{table}")
        ok = eq_actual >= eq_expected and fig_actual >= fig_expected and tab_actual >= tab_expected
        return ok, (
            f"equations {eq_actual}/{eq_expected}, figures {fig_actual}/{fig_expected}, "
            f"tables {tab_actual}/{tab_expected}"
        )

    def bib_present_and_nonempty(self, c):
        rel = c.get("target") or "05-template/build/references.bib"
        path = os.path.join(self.workdir, rel)
        if not os.path.isfile(path):
            return False, f"bib missing: {rel}"
        size = os.path.getsize(path)
        if size == 0 and not self.rewritten.get("references"):
            return None, "bibliography empty because rewritten IR has no references"
        return size > 0, f"bib bytes={size}"

    def template_source_official(self, c):
        official = self.manifest.get("is_official_template")
        if official is None:
            return None, "manifest missing is_official_template"
        return bool(official), f"is_official_template={official}"

    def zero_errors(self, c):
        n = (self.validation.get("summary") or {}).get("error_count")
        if not isinstance(n, int):
            return False, "validation.summary.error_count missing"
        return n == 0, f"error_count={n}"

    def citation_bidirectional_clean(self, c):
        findings = []
        for checker in self.validation.get("checkers", []):
            if checker.get("name") == "check_citations":
                findings.extend(checker.get("findings") or [])
        bad = [f.get("id") for f in findings if str(f.get("id", "")).startswith(("cite-", "bib-"))]
        return not bad, f"citation-direction findings={bad}"

    def numeric_consistency_clean(self, c):
        return self.numeric_consistency(c)

    def no_cjk_residue(self, c):
        checker = next(
            (x for x in self.validation.get("checkers", []) if x.get("name") == "check_cjk_residue"),
            None,
        )
        if checker is None:
            return None, "CJK checker not present"
        return checker.get("status") in ("pass", "skipped"), f"checker status={checker.get('status')}"

    def no_placeholder_residue(self, c):
        placeholders = common.find_placeholders(_all_text(self.rewritten))
        return not placeholders, f"placeholders={placeholders}"

    def latex_compiles(self, c):
        main = os.path.join(self.workdir, "05-template", "build", "main.tex")
        if not os.path.isfile(main):
            return False, "main.tex missing"
        latexmk = shutil.which("latexmk")
        pdflatex = shutil.which("pdflatex")
        if not (latexmk or pdflatex):
            return None, "latex toolchain unavailable (configured as skippable/degrade)"
        cmd = (
            [latexmk, "-pdf", "-interaction=nonstopmode", "-halt-on-error", "main.tex"]
            if latexmk else
            [pdflatex, "-interaction=nonstopmode", "-halt-on-error", "main.tex"]
        )
        proc = subprocess.run(
            cmd,
            cwd=os.path.dirname(main),
            capture_output=True,
            text=True,
            timeout=120,
        )
        return proc.returncode == 0, f"compile returncode={proc.returncode}"

    def journal_constraints_satisfied(self, c):
        findings = []
        for checker in self.validation.get("checkers", []):
            if checker.get("name") == "check_journal_constraints":
                findings.extend(checker.get("findings") or [])
        errors = [f.get("id") for f in findings if f.get("severity") == "error"]
        return not errors, f"constraint errors={errors}"

    def no_orphan_numbered_objects(self, c):
        findings = []
        for checker in self.validation.get("checkers", []):
            if checker.get("name") in ("check_figures_tables", "check_equations"):
                findings.extend(checker.get("findings") or [])
        orphans = [f.get("id") for f in findings if f.get("severity") in ("warning", "error")]
        return not orphans, f"orphan findings={orphans}"

    def language_quality(self, c):
        return None, "requires agent/human language review; not silently passed"


def evaluate_gate(gate_id: str, workdir: str, source_path: str | None = None) -> dict[str, Any]:
    return GateEvaluator(workdir, source_path=source_path).evaluate(gate_id)


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(description="Evaluate one configured quality gate")
    parser.add_argument("gate", choices=[f"G{i}" for i in range(1, 7)])
    parser.add_argument("--workdir", default="runs/demo")
    parser.add_argument("--source", default=None)
    args = parser.parse_args()
    result = evaluate_gate(args.gate, args.workdir, args.source)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    raise SystemExit(0 if result["passed"] else 2)
