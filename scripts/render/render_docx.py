#!/usr/bin/env python3
"""Compatibility CLI: render manuscript IR to a reviewable Word document.

Scientific content is copied from the IR. Equations are emitted as LaTeX text
rather than being silently converted to potentially lossy Word equations.
"""
from __future__ import annotations

import argparse
import os
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import common  # noqa: E402


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--ir", required=True)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()

    from docx import Document

    ms = common.load_json(args.ir)
    doc = Document()
    meta = ms.get("meta", {})
    doc.add_heading(meta.get("title") or "Untitled Manuscript", 0)
    if meta.get("abstract"):
        doc.add_heading("Abstract", level=1)
        doc.add_paragraph(meta["abstract"])
    if meta.get("keywords"):
        doc.add_paragraph("Keywords: " + ", ".join(meta["keywords"]))

    for sec in ms.get("sections", []):
        level = max(1, min(9, int(sec.get("level", 1))))
        doc.add_heading(sec.get("heading") or "Untitled section", level=level)
        for block in sec.get("blocks", []):
            kind = block.get("type")
            if kind == "paragraph":
                doc.add_paragraph(block.get("text") or "")
            elif kind == "equation":
                p = doc.add_paragraph()
                p.add_run(block.get("latex") or "")
            elif kind == "list":
                style = "List Number" if block.get("ordered") else "List Bullet"
                for item in block.get("items") or []:
                    doc.add_paragraph(str(item), style=style)
            elif kind in ("figure", "table"):
                doc.add_paragraph(f"[{kind.upper()} {block.get('ref_id') or ''} preserved in IR]")
            elif block.get("text"):
                doc.add_paragraph(block["text"])

    if ms.get("tables"):
        doc.add_heading("Tables", level=1)
        for tab in ms["tables"]:
            doc.add_paragraph(tab.get("caption") or tab.get("id") or "Table")
            header = tab.get("header") or []
            rows = tab.get("rows") or []
            ncols = max(1, len(header), max((len(r) for r in rows), default=0))
            t = doc.add_table(rows=1 if header else 0, cols=ncols)
            if header:
                for idx, value in enumerate(header):
                    t.rows[0].cells[idx].text = str(value)
            for row in rows:
                cells = t.add_row().cells
                for idx, value in enumerate(row[:ncols]):
                    cells[idx].text = str(value)

    os.makedirs(os.path.dirname(args.out) or ".", exist_ok=True)
    doc.save(args.out)
    print(args.out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
